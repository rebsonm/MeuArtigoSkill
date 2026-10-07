#!/usr/bin/env python3
"""Conservatively consolidate bibliographic exports.

Exact DOI/title duplicates are consolidated. Fuzzy title matches are only flagged
for manual/agent audit and are never removed automatically.

Supported input: CSV, TSV/TXT, XLSX (XLSX requires openpyxl).
"""

from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path

ALIASES = {
    "title": ["title", "article title", "document title", "ti"],
    "doi": ["doi", "digital object identifier", "di"],
    "authors": ["authors", "author", "author full names", "au"],
    "year": ["year", "publication year", "py"],
}
DOI_RE = re.compile(r"^10\.\d{4,9}/\S+$", re.I)


def norm_header(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def pick(row: dict[str, str], key: str) -> str:
    mapped = {norm_header(k): (v or "") for k, v in row.items()}
    for alias in ALIASES[key]:
        if alias in mapped and str(mapped[alias]).strip():
            return str(mapped[alias]).strip()
    return ""


def norm_doi(value: str) -> str:
    s = str(value or "").strip().lower()
    s = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", s)
    s = re.sub(r"^doi:\s*", "", s)
    s = s.strip().rstrip(".,;)")
    return s if DOI_RE.match(s) else ""


def norm_title(value: str) -> str:
    s = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def first_author(value: str) -> str:
    s = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    s = re.split(r";|\band\b|\|", s)[0]
    return re.sub(r"[^a-z]+", "", s)


def read_tabular(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise SystemExit("XLSX input requires openpyxl; convert to CSV or install openpyxl.") from e
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(x or "") for x in next(rows)]
        return [{headers[i]: "" if v is None else str(v) for i, v in enumerate(vals)} for vals in rows]

    delimiter = "\t" if suffix in {".tsv", ".txt"} else ","
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096); f.seek(0)
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
        except csv.Error:
            pass
        return list(csv.DictReader(f, delimiter=delimiter))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+", help="Bibliographic CSV/TSV/XLSX exports")
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--fuzzy-threshold", type=float, default=0.94)
    args = ap.parse_args()

    records = []
    all_fields = set()
    for f in args.inputs:
        p = Path(f).resolve()
        for idx, row in enumerate(read_tabular(p), start=2):
            row = {str(k): "" if v is None else str(v) for k, v in row.items()}
            row["_source_file"] = p.name
            row["_source_row"] = str(idx)
            row["_norm_doi"] = norm_doi(pick(row, "doi"))
            row["_norm_title"] = norm_title(pick(row, "title"))
            row["_first_author"] = first_author(pick(row, "authors"))
            row["_year"] = pick(row, "year")
            records.append(row)
            all_fields.update(row.keys())

    canonical = []
    exact_audit = []
    doi_index = {}
    title_index = {}

    for rec in records:
        match = None; reason = ""
        if rec["_norm_doi"] and rec["_norm_doi"] in doi_index:
            match = doi_index[rec["_norm_doi"]]; reason = "exact_doi"
        elif rec["_norm_title"] and rec["_norm_title"] in title_index:
            match = title_index[rec["_norm_title"]]; reason = "exact_title"

        if match is None:
            ci = len(canonical)
            canonical.append(rec)
            if rec["_norm_doi"]: doi_index[rec["_norm_doi"]] = ci
            if rec["_norm_title"]: title_index[rec["_norm_title"]] = ci
        else:
            exact_audit.append({
                "duplicate_source_file": rec["_source_file"],
                "duplicate_source_row": rec["_source_row"],
                "canonical_source_file": canonical[match]["_source_file"],
                "canonical_source_row": canonical[match]["_source_row"],
                "reason": reason,
                "doi": rec["_norm_doi"],
                "title": pick(rec, "title"),
            })

    fuzzy = []
    for i in range(len(canonical)):
        a = canonical[i]
        ta = a["_norm_title"]
        if not ta: continue
        for j in range(i + 1, len(canonical)):
            b = canonical[j]
            tb = b["_norm_title"]
            if not tb or ta == tb: continue
            author_ok = a["_first_author"] and a["_first_author"] == b["_first_author"]
            year_ok = a["_year"] and b["_year"] and a["_year"] == b["_year"]
            if not (author_ok or year_ok): continue
            score = SequenceMatcher(None, ta, tb).ratio()
            if score >= args.fuzzy_threshold:
                fuzzy.append({
                    "source_file_a": a["_source_file"], "source_row_a": a["_source_row"],
                    "title_a": pick(a, "title"), "doi_a": a["_norm_doi"],
                    "source_file_b": b["_source_file"], "source_row_b": b["_source_row"],
                    "title_b": pick(b, "title"), "doi_b": b["_norm_doi"],
                    "title_similarity": f"{score:.4f}",
                    "author_overlap_signal": "yes" if author_ok else "no",
                    "same_year_signal": "yes" if year_ok else "no",
                    "decision": "REVIEW_REQUIRED",
                    "notes": "Do not remove automatically.",
                })

    out = Path(args.out_dir).resolve(); out.mkdir(parents=True, exist_ok=True)
    fields = sorted(all_fields)
    with (out / "consolidated_exact_dedup.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(canonical)
    with (out / "exact_duplicate_audit.csv").open("w", newline="", encoding="utf-8-sig") as f:
        fields2 = ["duplicate_source_file","duplicate_source_row","canonical_source_file","canonical_source_row","reason","doi","title"]
        w=csv.DictWriter(f, fieldnames=fields2); w.writeheader(); w.writerows(exact_audit)
    with (out / "near_duplicate_review.csv").open("w", newline="", encoding="utf-8-sig") as f:
        fields3 = ["source_file_a","source_row_a","title_a","doi_a","source_file_b","source_row_b","title_b","doi_b","title_similarity","author_overlap_signal","same_year_signal","decision","notes"]
        w=csv.DictWriter(f, fieldnames=fields3); w.writeheader(); w.writerows(fuzzy)

    print(f"input_records={len(records)}")
    print(f"canonical_after_exact_dedup={len(canonical)}")
    print(f"exact_duplicates={len(exact_audit)}")
    print(f"near_duplicate_pairs_to_review={len(fuzzy)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
